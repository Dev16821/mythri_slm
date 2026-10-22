
#!/usr/bin/env python3

import json
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


HOME = Path.home()
REPORT_DIR = HOME / "codeql-databases" / "mythri-audit-reports"
OUTPUT_DIR = REPORT_DIR / "excel"
TIMESTAMP = datetime.now().strftime("%Y%m%d_%H%M%S")

SARIF_FILES = [
    REPORT_DIR / "mythri-security-results.sarif",
    REPORT_DIR / "mythri-quality-extended.sarif",
    REPORT_DIR / "mythri-security-extended.sarif",
]

# Also support the filenames produced by the audit runner.
SARIF_FILES += [
    REPORT_DIR / "python-security-and-quality.sarif",
    REPORT_DIR / "python-code-quality-extended.sarif",
    REPORT_DIR / "python-security-extended.sarif",
]

HEADERS = [
    "Finding ID",
    "Rule ID",
    "Description",
    "Message",
    "Level",
    "Security Score",
    "Source File",
    "Line",
    "Column",
    "Suite / Report",
    "Tags",
    "Duplicate Rule + Location",
]


def read_sarif(path):
    """Return normalized finding records from one SARIF report."""
    records = []
    data = json.loads(path.read_text(encoding="utf-8"))

    for run in data.get("runs", []):
        driver = run.get("tool", {}).get("driver", {})
        rules = {
            rule.get("id"): rule
            for rule in driver.get("rules", [])
        }

        for result in run.get("results", []):
            rule_id = result.get("ruleId", "unknown")
            rule = rules.get(rule_id, {})
            properties = rule.get("properties", {})

            tags = properties.get("tags", [])
            if not isinstance(tags, list):
                tags = []

            score = properties.get("security-severity")
            try:
                score = float(score) if score is not None else None
            except (TypeError, ValueError):
                score = None

            locations = result.get("locations", []) or [None]

            for location in locations:
                physical = (
                    location.get("physicalLocation", {})
                    if location else {}
                )
                artifact = physical.get("artifactLocation", {})
                region = physical.get("region", {})

                records.append({
                    "Finding ID": result.get("guid")
                        or result.get("correlationGuid")
                        or "",
                    "Rule ID": rule_id,
                    "Description": rule.get(
                        "shortDescription", {}
                    ).get("text", ""),
                    "Message": result.get("message", {}).get("text", ""),
                    "Level": result.get("level", "unspecified"),
                    "Security Score": score,
                    "Source File": artifact.get("uri", ""),
                    "Line": region.get("startLine"),
                    "Column": region.get("startColumn"),
                    "Suite / Report": path.name,
                    "Tags": ", ".join(str(tag) for tag in tags),
                    "Duplicate Rule + Location": "",
                })

    return records


def style_sheet(ws, table_name):
    """Apply filters, readable widths, frozen headers and table styling."""
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False

    if ws.max_row >= 2 and ws.max_column >= 1:
        ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
        table = Table(displayName=table_name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(table)

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="243B53")
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for column in ws.columns:
        letter = get_column_letter(column[0].column)
        values = [len(str(c.value or "")) for c in column[:100]]
        width = min(max(max(values, default=10) + 2, 12), 55)
        ws.column_dimensions[letter].width = width

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=cell.column in (3, 4, 7, 10, 11),
            )


def main():
    existing = list(dict.fromkeys(p for p in SARIF_FILES if p.is_file()))

    if not existing:
        raise FileNotFoundError(
            f"No SARIF files found in {REPORT_DIR}\n"
            "Check the actual filenames in that directory."
        )

    findings = []
    for path in existing:
        findings.extend(read_sarif(path))

    # Mark repeated rule/location pairs across suites; retain all rows.
    occurrence_counts = Counter(
        (r["Rule ID"], r["Source File"], r["Line"])
        for r in findings
    )

    for record in findings:
        key = (
            record["Rule ID"],
            record["Source File"],
            record["Line"],
        )
        if occurrence_counts[key] > 1:
            record["Duplicate Rule + Location"] = (
                f"Repeated {occurrence_counts[key]} times across reports"
            )

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    output_file = OUTPUT_DIR / f"mythri_codeql_{TIMESTAMP}.xlsx"

    wb = Workbook()
    all_ws = wb.active
    all_ws.title = "All Findings"
    all_ws.append(HEADERS)

    for record in findings:
        all_ws.append([record.get(header, "") for header in HEADERS])

    security_ws = wb.create_sheet("Security Findings")
    security_ws.append(HEADERS)

    for record in findings:
        # Security score is stronger evidence of security classification
        # than a suite filename alone.
        is_security = (
            record["Security Score"] is not None
            or "security" in record["Rule ID"].lower()
            or "external/cwe" in record["Tags"].lower()
        )
        if is_security:
            security_ws.append(
                [record.get(header, "") for header in HEADERS]
            )

    summary_ws = wb.create_sheet("Summary")
    summary_ws.append(["Metric", "Count"])
    summary_ws.append(["SARIF reports imported", len(existing)])
    summary_ws.append(["Total finding locations", len(findings)])
    summary_ws.append([
        "Locations with security-severity scores",
        sum(r["Security Score"] is not None for r in findings),
    ])
    summary_ws.append([
        "Repeated rule/location records",
        sum(bool(r["Duplicate Rule + Location"]) for r in findings),
    ])
    summary_ws.append([])
    summary_ws.append(["Rule ID", "Finding Count"])

    for rule, count in Counter(
        r["Rule ID"] for r in findings
    ).most_common():
        summary_ws.append([rule, count])

    summary_ws.append([])
    summary_ws.append(["Source File", "Finding Count"])
    for file, count in Counter(
        r["Source File"] or "(location unavailable)"
        for r in findings
    ).most_common():
        summary_ws.append([file, count])

    rules_ws = wb.create_sheet("Rule Reference")
    rules_ws.append(["Rule ID", "Description", "Report", "Tags"])
    seen_rules = set()

    for record in findings:
        key = (record["Rule ID"], record["Suite / Report"])
        if key in seen_rules:
            continue
        seen_rules.add(key)
        rules_ws.append([
            record["Rule ID"],
            record["Description"],
            record["Suite / Report"],
            record["Tags"],
        ])

    for idx, ws in enumerate(wb.worksheets, start=1):
        safe_name = f"AuditTable{idx}"
        style_sheet(ws, safe_name)

    wb.save(output_file)

    print(f"Imported reports : {len(existing)}")
    print(f"Finding locations: {len(findings)}")
    print(f"Excel workbook   : {output_file}")


if __name__ == "__main__":
    main()
