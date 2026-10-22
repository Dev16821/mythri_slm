
#!/usr/bin/env python3
"""
Mythri Data-Flow Audit
======================
Runs CodeQL queries against an existing Python database and exports
findings and any SARIF data-flow paths to an Excel workbook.

Default query suite:
  python-security-extended.qls

Optionally supply a custom .ql query or .qls suite designed for
Mythri-specific data-flow analysis.

Does not modify application source code.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import subprocess
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils import get_column_letter
except ImportError:
    print(
        "Missing dependency. Install with:\n"
        "python3 -m pip install --user openpyxl",
        file=sys.stderr,
    )
    raise SystemExit(2)


HOME = Path.home()

DEFAULT_CODEQL = (
    HOME
    / ".config/Code/User/globalStorage/github.vscode-codeql"
    / "distribution1/codeql/codeql"
)

DEFAULT_DATABASE = HOME / "codeql-databases/mythri-python"

DEFAULT_SUITE = (
    HOME
    / ".codeql/packages/codeql/python-queries/1.8.12"
    / "codeql-suites/python-security-extended.qls"
)

DEFAULT_OUTPUT = HOME / "codeql-databases/mythri-data-flow-reports"


FINDING_HEADERS = [
    "Severity",
    "Rule ID",
    "Rule Name",
    "Description",
    "File",
    "Line",
    "Column",
    "Finding Fingerprint",
    "Explicit Flow Path Available",
    "Analysis Source",
]

FLOW_HEADERS = [
    "Finding Fingerprint",
    "Rule ID",
    "Severity",
    "Flow Path Number",
    "Step Number",
    "Step Message",
    "File",
    "Line",
    "Column",
    "Code Snippet",
]

SUMMARY_HEADERS = [
    "Rule ID",
    "Rule Name",
    "Severity",
    "Finding Count",
    "Description",
]


def run_command(command: list[str]) -> None:
    """Run CodeQL and display its output."""
    print("\nExecuting:")
    print(" ".join(str(arg) for arg in command))

    result = subprocess.run(
        command,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
    )

    if result.stdout:
        print(result.stdout)

    if result.returncode != 0:
        raise RuntimeError(
            f"CodeQL exited with status {result.returncode}"
        )


def sarif_location(location: dict[str, Any]) -> dict[str, Any]:
    """Extract file, line, column, message, and snippet from a SARIF location."""
    physical = location.get("physicalLocation", {})
    artifact = physical.get("artifactLocation", {})
    region = physical.get("region", {})

    snippet = region.get("snippet", {})
    if isinstance(snippet, dict):
        snippet = snippet.get("text", "")

    message_obj = location.get("message", {})
    message = (
        message_obj.get("text", "")
        if isinstance(message_obj, dict)
        else str(message_obj)
    )

    return {
        "file": artifact.get("uri", ""),
        "line": region.get("startLine", ""),
        "column": region.get("startColumn", ""),
        "snippet": snippet or "",
        "message": message or "",
    }


def load_sarif(
    sarif_path: Path,
    analysis_source: str,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Convert SARIF findings and explicit code-flow paths into worksheet rows."""
    document = json.loads(sarif_path.read_text(encoding="utf-8"))

    findings: list[dict[str, Any]] = []
    flow_steps: list[dict[str, Any]] = []

    for run in document.get("runs", []):
        driver = run.get("tool", {}).get("driver", {})
        rules = {
            rule.get("id", ""): rule
            for rule in driver.get("rules", [])
        }

        for result in run.get("results", []):
            rule_id = result.get("ruleId", "")
            rule = rules.get(rule_id, {})

            description_obj = rule.get("shortDescription", {})
            description = (
                description_obj.get("text", "")
                if isinstance(description_obj, dict)
                else str(description_obj)
            )

            locations = result.get("locations", [])
            primary = (
                sarif_location(locations[0].get("physicalLocation", {}))
                if locations
                else {"file": "", "line": "", "column": ""}
            )

            # Prefer CodeQL's supplied partial fingerprint where available.
            partial = result.get("partialFingerprints", {})
            fingerprint = (
                partial.get("primaryLocationLineHash")
                or partial.get("primaryLocationFingerprint")
            )

            if not fingerprint:
                fingerprint_input = "|".join([
                    rule_id,
                    str(primary["file"]),
                    str(primary["line"]),
                    result.get("message", {}).get("text", ""),
                ])
                fingerprint = hashlib.sha256(
                    fingerprint_input.encode("utf-8")
                ).hexdigest()[:24]

            severity = result.get("level") or (
                rule.get("defaultConfiguration", {}).get("level")
                or "unspecified"
            )

            code_flows = result.get("codeFlows", [])
            has_flow = bool(code_flows)

            findings.append({
                "Severity": str(severity).lower(),
                "Rule ID": rule_id,
                "Rule Name": rule.get("name", rule_id),
                "Description": (
                    result.get("message", {}).get("text", "")
                    or description
                ),
                "File": primary["file"],
                "Line": primary["line"],
                "Column": primary["column"],
                "Finding Fingerprint": fingerprint,
                "Explicit Flow Path Available": (
                    "Yes" if has_flow else "No path in SARIF"
                ),
                "Analysis Source": analysis_source,
            })

            # A codeFlow is a sequence of threadFlows; each threadFlow
            # contains locations that CodeQL associates with the flow.
            for path_number, code_flow in enumerate(code_flows, start=1):
                for thread_flow in code_flow.get("threadFlows", []):
                    for step_number, item in enumerate(
                        thread_flow.get("locations", []),
                        start=1,
                    ):
                        location = item.get("location", {})
                        extracted = sarif_location(location)
                        message_obj = location.get("message", {})
                        step_message = (
                            message_obj.get("text", "")
                            if isinstance(message_obj, dict)
                            else str(message_obj)
                        )

                        flow_steps.append({
                            "Finding Fingerprint": fingerprint,
                            "Rule ID": rule_id,
                            "Severity": str(severity).lower(),
                            "Flow Path Number": path_number,
                            "Step Number": step_number,
                            "Step Message": (
                                step_message
                                or extracted["message"]
                                or "No step description supplied"
                            ),
                            "File": extracted["file"],
                            "Line": extracted["line"],
                            "Column": extracted["column"],
                            "Code Snippet": extracted["snippet"],
                        })

    return findings, flow_steps


def style_worksheet(ws) -> None:
    """Apply filters, frozen headers, readable widths, and table styling."""
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False

    if ws.max_row == 0 or ws.max_column == 0:
        return

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="244062")
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    ws.row_dimensions[1].height = 32
    ws.auto_filter.ref = ws.dimensions

    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        header = str(ws.cell(1, col_idx).value or "")
        sample_values = [
            str(ws.cell(row, col_idx).value or "")
            for row in range(1, min(ws.max_row, 80) + 1)
        ]

        width = min(
            max(max((len(value) for value in sample_values), default=12) + 2, 12),
            48,
        )

        if header in {
            "Description", "Step Message", "Code Snippet", "File",
            "Analysis Source",
        }:
            width = 48

        ws.column_dimensions[letter].width = width

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    if ws.max_row > 1:
        safe_title = "".join(
            char for char in ws.title if char.isalnum()
        )[:20] or "Results"

        table = Table(
            displayName=f"T{safe_title}",
            ref=ws.dimensions,
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showRowStripes=True,
            showColumnStripes=False,
            showFirstColumn=False,
            showLastColumn=False,
        )
        ws.add_table(table)


def add_sheet(
    workbook: Workbook,
    title: str,
    headers: list[str],
    rows: list[list[Any]],
) -> None:
    ws = workbook.create_sheet(title)
    ws.append(headers)

    for row in rows:
        ws.append(row)

    style_worksheet(ws)


def create_excel_report(
    output_path: Path,
    findings: list[dict[str, Any]],
    flow_steps: list[dict[str, Any]],
    execution_log: list[dict[str, str]],
) -> None:
    """Create the workbook and its summary sheets."""
    workbook = Workbook()

    summary = workbook.active
    summary.title = "Executive Summary"

    summary.append(["Metric", "Value"])
    summary.append([
        "Generated At",
        datetime.now().astimezone().isoformat(),
    ])
    summary.append(["Total Findings", len(findings)])
    summary.append(["Explicit Flow Steps", len(flow_steps)])
    summary.append([
        "Findings with Flow Paths",
        sum(
            1 for item in findings
            if item["Explicit Flow Path Available"] == "Yes"
        ),
    ])
    summary.append([
        "Findings without Flow Paths in SARIF",
        sum(
            1 for item in findings
            if item["Explicit Flow Path Available"] != "Yes"
        ),
    ])
    summary.append([])
    summary.append(["Severity", "Count"])

    for severity, count in sorted(
        Counter(item["Severity"] for item in findings).items()
    ):
        summary.append([severity, count])

    summary.append([])
    summary.append([
        "Interpretation",
        "A finding is a query result, not automatically a confirmed defect.",
    ])
    summary.append([
        "Flow paths",
        "Only paths supplied by CodeQL in SARIF appear in the Flow Paths sheet.",
    ])
    summary.append([
        "Coverage",
        "The configured suite determines which flows are analyzed.",
    ])

    style_worksheet(summary)

    add_sheet(
        workbook,
        "Data Flow Findings",
        FINDING_HEADERS,
        [
            [item.get(header, "") for header in FINDING_HEADERS]
            for item in findings
        ],
    )

    add_sheet(
        workbook,
        "Flow Paths",
        FLOW_HEADERS,
        [
            [item.get(header, "") for header in FLOW_HEADERS]
            for item in flow_steps
        ],
    )

    grouped: dict[str, list[dict[str, Any]]] = {}
    for item in findings:
        grouped.setdefault(item["Rule ID"], []).append(item)

    rule_rows = []
    for rule_id, items in sorted(grouped.items()):
        first = items[0]
        rule_rows.append([
            rule_id,
            first["Rule Name"],
            first["Severity"],
            len(items),
            first["Description"],
        ])

    add_sheet(
        workbook,
        "Rule Summary",
        SUMMARY_HEADERS,
        rule_rows,
    )

    add_sheet(
        workbook,
        "Execution Log",
        ["Analysis", "Status", "Query Path", "Details"],
        [
            [
                item["analysis"],
                item["status"],
                item["query_path"],
                item["details"],
            ]
            for item in execution_log
        ],
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Run CodeQL data-flow analysis and export to Excel."
    )

    parser.add_argument("--codeql", type=Path, default=DEFAULT_CODEQL)
    parser.add_argument("--database", type=Path, default=DEFAULT_DATABASE)
    parser.add_argument("--suite", type=Path, default=DEFAULT_SUITE)
    parser.add_argument(
        "--custom-query",
        type=Path,
        help="Optional custom Mythri-specific .ql query or .qls suite.",
    )
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)

    args = parser.parse_args()

    for label, path in [
        ("CodeQL CLI", args.codeql),
        ("CodeQL database", args.database),
        ("Default query suite", args.suite),
    ]:
        if not path.exists():
            print(f"ERROR: {label} not found: {path}", file=sys.stderr)
            return 2

    if args.custom_query and not args.custom_query.exists():
        print(
            f"ERROR: Custom query not found: {args.custom_query}",
            file=sys.stderr,
        )
        return 2

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    run_dir = args.output_dir / timestamp
    run_dir.mkdir(parents=True, exist_ok=False)

    all_findings: list[dict[str, Any]] = []
    all_flow_steps: list[dict[str, Any]] = []
    execution_log: list[dict[str, str]] = []

    queries = [("Extended security suite", args.suite)]

    if args.custom_query:
        queries.append(("Custom Mythri query", args.custom_query))

    try:
        for index, (label, query_path) in enumerate(queries, start=1):
            sarif_path = run_dir / f"data_flow_{index}.sarif"

            try:
                run_command([
                    str(args.codeql),
                    "database", "analyze",
                    str(args.database),
                    str(query_path),
                    "--format=sarif-latest",
                    f"--output={sarif_path}",
                    "--threads=0",
                ])

                findings, steps = load_sarif(sarif_path, label)
                all_findings.extend(findings)
                all_flow_steps.extend(steps)

                execution_log.append({
                    "analysis": label,
                    "status": "Success",
                    "query_path": str(query_path),
                    "details": (
                        f"{len(findings)} findings; "
                        f"{len(steps)} explicit flow steps; "
                        f"SARIF: {sarif_path.name}"
                    ),
                })

            except Exception as exc:
                execution_log.append({
                    "analysis": label,
                    "status": "Failed",
                    "query_path": str(query_path),
                    "details": str(exc),
                })
                print(f"WARNING: {label} failed: {exc}", file=sys.stderr)

        # Deduplicate identical findings produced by overlapping queries
        # from the same analysis source.
        unique_findings = {}
        for item in all_findings:
            key = (
                item["Analysis Source"],
                item["Rule ID"],
                item["Finding Fingerprint"],
            )
            unique_findings[key] = item

        all_findings = list(unique_findings.values())

        workbook_path = (
            run_dir / f"mythri_data_flow_audit_{timestamp}.xlsx"
        )

        create_excel_report(
            workbook_path,
            all_findings,
            all_flow_steps,
            execution_log,
        )

        print("\nDATA-FLOW AUDIT FINISHED")
        print(f"Findings: {len(all_findings)}")
        print(f"Explicit flow steps: {len(all_flow_steps)}")
        print(f"Excel workbook: {workbook_path}")
        print(f"Raw SARIF files: {run_dir}")

        if not any(item["status"] == "Success" for item in execution_log):
            print(
                "No query completed successfully. Inspect the Execution Log "
                "and CodeQL output.",
                file=sys.stderr,
            )
            return 1

        return 0

    except Exception as exc:
        print(f"AUDIT FAILED: {exc}", file=sys.stderr)
        print(f"Partial results, if any: {run_dir}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
