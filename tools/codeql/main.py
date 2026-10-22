#!/usr/bin/env python3
"""
Mythri Master Audit Runner
==========================
Runs every Mythri audit script in order and merges the results into ONE
Excel workbook with several readable sheets.

Place this file next to the other audit scripts (<repo>/codeql/main.py) or in
the repository root, then run:

    python3 codeql/main.py

Useful options:
    --skip codeql                 reuse the existing CodeQL database (faster)
    --skip data_flow control_flow skip stages you do not need this time
    --quality-skip-codeql         run only the AST checks inside quality_audit
    --with-custom-queries         also run the custom .ql templates (see notes)
    --baseline HEAD~3             compare the Git change audit to another rev

Stages (each is an existing script, run unchanged):
    codeql        codeql_audit.py          rebuild DB + 3 query suites
    data_flow     data_flow_audit.py       extended security suite + flow paths
    control_flow  control_flow_audit.py    quality suite, keyword triage
    quality       quality_audit.py         AST metrics (+ CodeQL quality suite)
    changes       mythri_change_audit.py   read-only Git change audit

No application source file is modified by this script.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import shlex
import shutil
import subprocess
import sys
import time
from collections import Counter, defaultdict, deque
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit(
        "Missing dependency. Install with:\n"
        "python3 -m pip install --user openpyxl"
    )


# --------------------------------------------------------------------------
# Locations
# --------------------------------------------------------------------------

HOME = Path.home()
HERE = Path(__file__).resolve().parent

if (HERE / "codeql_audit.py").is_file():
    SCRIPTS_DIR = HERE
    DEFAULT_REPO = HERE.parent
else:
    SCRIPTS_DIR = HERE / "codeql"
    DEFAULT_REPO = HERE

DB_PATH = HOME / "codeql-databases" / "mythri-python"
LEGACY_REPORT_DIR = HOME / "codeql-databases" / "mythri-audit-reports"
DEFAULT_OUT = HOME / "codeql-databases" / "mythri-master-reports"

LEGACY_SUITES = (
    "python-security-and-quality",
    "python-code-quality-extended",
    "python-security-extended",
)

STAGE_KEYS = ("codeql", "data_flow", "control_flow", "quality", "changes")

ENV = {**os.environ, "PYTHONUNBUFFERED": "1"}


# --------------------------------------------------------------------------
# Tool discovery
# --------------------------------------------------------------------------

def find_codeql() -> Path | None:
    root = HOME / ".config/Code/User/globalStorage/github.vscode-codeql"
    candidates = [root / "distribution1/codeql/codeql"]
    on_path = shutil.which("codeql")
    if on_path:
        candidates.append(Path(on_path))
    if root.is_dir():
        candidates.extend(sorted(root.glob("distribution*/codeql/codeql"),
                                 reverse=True))
    for candidate in candidates:
        if candidate.is_file() and os.access(candidate, os.X_OK):
            return candidate
    return None


def _version_key(path: Path) -> tuple[int, ...]:
    return tuple(
        int(part) if part.isdigit() else 0
        for part in re.split(r"[.\-]", path.name)
    )


def find_suite_dir() -> Path | None:
    """Newest installed python-queries pack (numeric version order)."""
    root = HOME / ".codeql/packages/codeql/python-queries"
    if not root.is_dir():
        return None
    versions = sorted(
        (p for p in root.iterdir() if p.is_dir()),
        key=_version_key,
        reverse=True,
    )
    for version in versions:
        suites = version / "codeql-suites"
        if ((suites / "python-security-extended.qls").is_file()
                and (suites / "python-code-quality-extended.qls").is_file()):
            return suites
    return None


# --------------------------------------------------------------------------
# Running stages
# --------------------------------------------------------------------------

@dataclass
class StageResult:
    key: str
    title: str
    status: str = "NOT RUN"
    seconds: float = 0.0
    command: str = ""
    detail: str = ""
    tail: list[str] = field(default_factory=list)


def run_stage(key: str, title: str, cmd: list[Any], cwd: Path) -> StageResult:
    result = StageResult(key, title)
    result.command = " ".join(shlex.quote(str(part)) for part in cmd)

    print(f"\n{'#' * 76}\n# {title}\n{'#' * 76}", flush=True)
    print("$", result.command, flush=True)

    started = time.monotonic()
    tail: deque[str] = deque(maxlen=40)

    try:
        proc = subprocess.Popen(
            [str(part) for part in cmd],
            cwd=str(cwd),
            env=ENV,
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=True,
            errors="replace",
            bufsize=1,
        )
        assert proc.stdout is not None
        for line in proc.stdout:
            sys.stdout.write(line)
            tail.append(line.rstrip())
        code = proc.wait()
        result.status = "SUCCESS" if code == 0 else "FAILED"
        result.detail = f"exit code {code}"
    except OSError as exc:
        result.status = "FAILED"
        result.detail = f"could not start: {exc}"

    result.seconds = round(time.monotonic() - started, 1)
    result.tail = list(tail)
    print(f"\n[{result.status}] {title} ({result.seconds}s)", flush=True)
    return result


def skipped(key: str, title: str, reason: str) -> StageResult:
    print(f"\n[SKIPPED] {title}: {reason}", flush=True)
    return StageResult(key, title, status="SKIPPED", detail=reason)


# --------------------------------------------------------------------------
# SARIF parsing
# --------------------------------------------------------------------------

SEV_RANK = {"CRITICAL": 4, "HIGH": 3, "MEDIUM": 2, "LOW": 1, "INFO": 0}


def severity_of(level: str, score: float | None) -> str:
    if score is not None:
        if score >= 9:
            return "CRITICAL"
        if score >= 7:
            return "HIGH"
        if score >= 4:
            return "MEDIUM"
        return "LOW"
    return {
        "error": "HIGH",
        "warning": "MEDIUM",
        "note": "LOW",
        "recommendation": "LOW",
    }.get((level or "").lower(), "INFO")


def norm_path(path: Any) -> str:
    text = str(path or "").replace("\\", "/")
    return text[2:] if text.startswith("./") else text


def location_of(location: dict[str, Any] | None) -> dict[str, Any]:
    physical = (location or {}).get("physicalLocation", {}) or {}
    artifact = physical.get("artifactLocation", {}) or {}
    region = physical.get("region", {}) or {}
    snippet = region.get("snippet", {})
    if isinstance(snippet, dict):
        snippet = snippet.get("text", "")
    return {
        "file": norm_path(artifact.get("uri", "")),
        "line": region.get("startLine"),
        "snippet": snippet or "",
    }


def parse_sarif(path: Path, source: str) -> list[dict[str, Any]]:
    """Convert one SARIF file into flat finding records (with flow steps)."""
    document = json.loads(path.read_text(encoding="utf-8"))
    findings: list[dict[str, Any]] = []

    for run in document.get("runs", []):
        driver = run.get("tool", {}).get("driver", {})
        rules = {r.get("id"): r for r in driver.get("rules", [])}

        for res in run.get("results", []):
            rule_id = res.get("ruleId") or "unknown"
            rule = rules.get(rule_id, {})
            props = rule.get("properties") or {}

            try:
                score = float(props.get("security-severity"))
            except (TypeError, ValueError):
                score = None

            level = (
                res.get("level")
                or (rule.get("defaultConfiguration") or {}).get("level", "")
            )
            locations = res.get("locations") or []
            primary = location_of(locations[0] if locations else None)

            flows: list[dict[str, Any]] = []
            for path_no, code_flow in enumerate(res.get("codeFlows") or [], 1):
                for thread in code_flow.get("threadFlows") or []:
                    for step_no, item in enumerate(
                            thread.get("locations") or [], 1):
                        loc = item.get("location") or {}
                        where = location_of(loc)
                        flows.append({
                            "path_no": path_no,
                            "step_no": step_no,
                            "message": (loc.get("message") or {}).get(
                                "text", ""),
                            **where,
                        })

            findings.append({
                "rule_id": rule_id,
                "rule_name": rule.get("name") or rule_id,
                "description": (
                    rule.get("shortDescription") or {}).get("text", ""),
                "message": (res.get("message") or {}).get("text", ""),
                "severity": severity_of(level, score),
                "score": score,
                "tags": [str(t) for t in (props.get("tags") or [])],
                "file": primary["file"],
                "line": primary["line"],
                "flows": flows,
                "source": source,
            })

    return findings


def merge_findings(findings: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """One row per unique rule+location+message; remember every reporter."""
    merged: dict[tuple, dict[str, Any]] = {}

    for item in findings:
        key = (item["rule_id"], item["file"], item["line"], item["message"])
        if key not in merged:
            item = dict(item)
            item["sources"] = {item["source"]}
            merged[key] = item
            continue
        kept = merged[key]
        kept["sources"].add(item["source"])
        if not kept["flows"] and item["flows"]:
            kept["flows"] = item["flows"]
        if SEV_RANK[item["severity"]] > SEV_RANK[kept["severity"]]:
            kept["severity"] = item["severity"]
            kept["score"] = item["score"]

    ordered = sorted(
        merged.values(),
        key=lambda f: (
            -SEV_RANK[f["severity"]],
            -(f["score"] or 0),
            f["file"],
            f["line"] or 0,
        ),
    )
    for number, item in enumerate(ordered, 1):
        item["id"] = f"F-{number:04d}"
    return ordered


def load_control_flow_rows(sarif_path: Path | None) -> tuple[list[list], str]:
    """Reuse control_flow_audit.py's own keyword triage."""
    if not sarif_path:
        return [], "control_flow stage produced no SARIF"
    sys.path.insert(0, str(SCRIPTS_DIR))
    try:
        import control_flow_audit as cfa
        results, meta = cfa.load_sarif(sarif_path)
    except Exception as exc:  # noqa: BLE001 - report, don't crash
        return [], f"could not read control-flow results: {exc}"

    rows = []
    for result in results:
        if not cfa.is_control_flow_candidate(result, meta):
            continue
        detail = cfa.finding_details(result, meta)
        searchable = (
            cfa.get_rule_text(detail["rule_id"], meta)
            + " " + detail["message"].lower()
        )
        matched = ", ".join(
            k for k in cfa.CONTROL_FLOW_KEYWORDS if k in searchable
        )
        rows.append([
            detail["severity"].upper() if detail["severity"] else "INFO",
            detail["rule_id"], detail["rule_name"], detail["file"],
            detail["line"], detail["message"], matched,
        ])
    return rows, ""


# --------------------------------------------------------------------------
# Excel helpers
# --------------------------------------------------------------------------

_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
THIN = Side(style="thin", color="D0D5DD")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill("solid", fgColor="1F3864")

STYLES = {
    "CRITICAL": ("C00000", "FFFFFF"),
    "HIGH": ("FFC7CE", "9C0006"),
    "MEDIUM": ("FFEB9C", "7F6000"),
    "LOW": ("DDEBF7", "1F4E78"),
    "INFO": ("F2F2F2", "404040"),
    "SUCCESS": ("C6EFCE", "006100"),
    "FAILED": ("FFC7CE", "9C0006"),
    "SKIPPED": ("EDEDED", "595959"),
    "NOT RUN": ("EDEDED", "595959"),
}


def clean(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (bool, int, float)):
        return value
    text = _ILLEGAL.sub("", str(value))
    if len(text) > 32000:
        text = text[:32000] + " ...[truncated]"
    return text


def put(ws, row: int, col: int, value: Any):
    cell = ws.cell(row=row, column=col)
    value = clean(value)
    cell.value = value
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"  # never let report text become a formula
    return cell


def style_cell(cell) -> None:
    style = STYLES.get(str(cell.value).strip().upper())
    if style:
        fill, font = style
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.font = Font(bold=True, color=font)
        cell.alignment = Alignment(horizontal="center", vertical="top")


def write_table(
    wb: Workbook,
    title: str,
    headers: list[str],
    rows: list[list],
    *,
    widths: dict[str, int] | None = None,
    wrap: tuple[str, ...] = (),
    styled: tuple[str, ...] = (),
    empty_note: str = "No records.",
):
    ws = wb.create_sheet(title[:31])
    widths = widths or {}

    for col, header in enumerate(headers, 1):
        cell = put(ws, 1, col, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[1].height = 30

    rows = list(rows) or [[empty_note] + [""] * (len(headers) - 1)]
    wrap_cols = {headers.index(h) + 1 for h in wrap if h in headers}
    styled_cols = {headers.index(h) + 1 for h in styled if h in headers}

    for r, row in enumerate(rows, 2):
        for c, value in enumerate(row, 1):
            cell = put(ws, r, c, value)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=c in wrap_cols)
            if c in styled_cols:
                style_cell(cell)

    for col, header in enumerate(headers, 1):
        if header in widths:
            width = widths[header]
        else:
            sample = [len(str(header))] + [
                len(str(row[col - 1])) for row in rows[:200]
                if col - 1 < len(row)
            ]
            width = min(max(max(sample) + 2, 10), 50)
        ws.column_dimensions[get_column_letter(col)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    return ws


def read_sheet(xlsx: Path | None, name: str) -> list[list]:
    """Read one sheet (header + rows) from another workbook."""
    if not xlsx or not xlsx.is_file():
        return []
    try:
        book = load_workbook(xlsx, read_only=True, data_only=True)
    except Exception:  # noqa: BLE001
        return []
    try:
        if name not in book.sheetnames:
            return []
        return [list(row) for row in book[name].iter_rows(values_only=True)]
    finally:
        book.close()


def as_dicts(rows: list[list]) -> list[dict]:
    if len(rows) < 2:
        return []
    headers = [str(h) for h in rows[0]]
    return [dict(zip(headers, row)) for row in rows[1:]]


# --------------------------------------------------------------------------
# Dashboard
# --------------------------------------------------------------------------

def dash_section(ws, row: int, text: str) -> int:
    cell = put(ws, row, 1, text)
    cell.font = Font(bold=True, size=13, color="1F3864")
    return row + 1


def dash_table(ws, row: int, headers: list[str], rows: list[list],
               styled: tuple[str, ...] = ()) -> int:
    for col, header in enumerate(headers, 1):
        cell = put(ws, row, col, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.border = BORDER
    for r, values in enumerate(rows, row + 1):
        for c, value in enumerate(values, 1):
            cell = put(ws, r, c, value)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if headers[c - 1] in styled:
                style_cell(cell)
    return row + len(rows) + 2


def build_dashboard(ws, meta, stage_results, merged, flow_count,
                    control_count, hotspots, git_summary, warnings) -> None:
    ws.sheet_view.showGridLines = False
    for letter, width in zip("ABCDEF", (36, 22, 16, 16, 16, 60)):
        ws.column_dimensions[letter].width = width

    title = put(ws, 1, 1, "Mythri SLM - Master Audit Report")
    title.font = Font(bold=True, size=18, color="1F3864")
    row = 3

    row = dash_section(ws, row, "Run information")
    row = dash_table(ws, row, ["Item", "Value"], [
        ["Generated", meta["generated"]],
        ["Repository", meta["repo"]],
        ["Raw output folder", meta["run_dir"]],
        ["CodeQL CLI", meta["cli"]],
        ["CodeQL database", meta["db"]],
        ["Database last modified", meta["db_mtime"]],
    ])

    row = dash_section(ws, row, "Stage results")
    row = dash_table(
        ws, row, ["Stage", "Status", "Seconds", "Details"],
        [[s.title, s.status, s.seconds, s.detail] for s in stage_results],
        styled=("Status",),
    )

    sev_counts = Counter(f["severity"] for f in merged)
    row = dash_section(ws, row, "Findings overview (duplicates across suites merged)")
    overview = [["Unique findings", len(merged)]]
    overview += [[sev, sev_counts.get(sev, 0)]
                 for sev in ("CRITICAL", "HIGH", "MEDIUM", "LOW", "INFO")]
    overview += [
        ["With CodeQL security score",
         sum(1 for f in merged if f["score"] is not None)],
        ["With an explicit data-flow path", flow_count],
        ["Control-flow candidates (keyword triage)", control_count],
        ["Files with findings", len({f["file"] for f in merged if f["file"]})],
    ]
    row = dash_table(ws, row, ["Metric", "Count"], overview, styled=("Metric",))

    top_rules = Counter(f["rule_id"] for f in merged).most_common(10)
    row = dash_section(ws, row, "Top 10 rules")
    row = dash_table(ws, row, ["Rule ID", "Findings"],
                     [[r, c] for r, c in top_rules] or [["None", 0]])

    row = dash_section(ws, row, "Top 10 hotspot files")
    row = dash_table(
        ws, row, ["File", "Findings", "High+Critical", "Git priority"],
        [[h[0], h[1], h[2] + h[3], h[13]] for h in hotspots[:10]]
        or [["None", 0, 0, ""]],
        styled=("Git priority",),
    )

    if git_summary:
        row = dash_section(ws, row, "Git change context")
        row = dash_table(ws, row, ["Item", "Value"], git_summary)

    if warnings:
        row = dash_section(ws, row, "Warnings")
        row = dash_table(ws, row, ["Warning"], [[w] for w in warnings])

    row = dash_section(ws, row, "Sheet guide")
    row = dash_table(ws, row, ["Sheet", "What it contains"], [
        ["Security Findings", "Findings with a CodeQL security score or CWE/security tag"],
        ["All Findings", "Every unique finding, newest suites merged, with 'Reported By'"],
        ["Rule Summary", "One row per rule: count, worst severity, files affected"],
        ["File Hotspots", "Per-file findings joined with AST metrics and Git change priority"],
        ["Data Flow Paths", "Step-by-step source-to-sink paths supplied by CodeQL"],
        ["Control Flow Candidates", "Findings matching control-flow keywords (triage aid)"],
        ["Git Changes", "Files changed vs the baseline, with path-based priority"],
        ["Custom Queries", "Status of the optional custom .ql templates"],
        ["AST Metrics / AST Review", "Per-file size metrics and heuristic review prompts"],
        ["Run Log", "Commands, durations and the last output lines of each stage"],
    ])

    row = dash_section(ws, row, "How to read this report")
    notes = [
        "A finding is a lead for investigation, not proof of a defect.",
        "No findings does not prove the code is defect-free.",
        "Severity comes from the CodeQL security score when present, "
        "otherwise from the SARIF level (error=HIGH, warning=MEDIUM, note=LOW).",
        "Git priority is a path-name heuristic (safety, policy, model, config...).",
    ]
    for note in notes:
        cell = put(ws, row, 1, note)
        cell.alignment = Alignment(wrap_text=False)
        row += 1


# --------------------------------------------------------------------------
# Whole workbook
# --------------------------------------------------------------------------

def build_report(out_path, meta, stage_results, parsed, control_rows,
                 quality_xlsx, change_xlsx, warnings) -> dict[str, int]:
    merged = merge_findings(parsed)

    ast_metrics = {
        norm_path(d["File"]): d
        for d in as_dicts(read_sheet(quality_xlsx, "File Metrics"))
    }
    ast_rows = read_sheet(quality_xlsx, "AST Review Candidates")
    ast_by_file = Counter(norm_path(d["File"]) for d in as_dicts(ast_rows))

    changed_rows = read_sheet(change_xlsx, "Changed Files")
    git_by_file = {
        norm_path(d["Path"]): d for d in as_dicts(changed_rows)
    }
    git_summary = [
        r[:2] for r in read_sheet(change_xlsx, "Executive Summary")[1:]
        if len(r) >= 2 and r[0]
    ]

    # ---- hotspots ----------------------------------------------------
    sev_by_file: dict[str, Counter] = defaultdict(Counter)
    rules_by_file: dict[str, Counter] = defaultdict(Counter)
    for f in merged:
        if f["file"]:
            sev_by_file[f["file"]][f["severity"]] += 1
            rules_by_file[f["file"]][f["rule_id"]] += 1

    files = set(sev_by_file) | set(ast_by_file) | {
        p for p, d in git_by_file.items()
        if d.get("Priority") in ("CRITICAL", "HIGH")
    }

    hotspots = []
    for path in files:
        sev = sev_by_file.get(path, Counter())
        ast = ast_metrics.get(path, {})
        git = git_by_file.get(path, {})
        top_rule = (rules_by_file[path].most_common(1)[0][0]
                    if rules_by_file.get(path) else "")
        hotspots.append([
            path, sum(sev.values()), sev["CRITICAL"], sev["HIGH"],
            sev["MEDIUM"], sev["LOW"], top_rule, ast_by_file.get(path, 0),
            ast.get("Lines", ""), ast.get("Functions", ""),
            ast.get("Long functions (>50 lines)", ""),
            ast.get("Bare except handlers", ""),
            git.get("Working Status", ""), git.get("Priority", ""),
        ])
    hotspots.sort(key=lambda h: (-(h[2] + h[3]), -h[1], -h[7], h[0]))

    # ---- flows -------------------------------------------------------
    flow_rows = []
    flow_findings = 0
    for f in merged:
        if not f["flows"]:
            continue
        flow_findings += 1
        alert = f"{f['file']}:{f['line']}"
        for step in f["flows"]:
            flow_rows.append([
                f["id"], f["rule_id"], f["severity"], alert,
                step["path_no"], step["step_no"],
                step["message"] or "(no step description)",
                step["file"], step["line"], step["snippet"],
            ])

    # ---- sheets ------------------------------------------------------
    wb = Workbook()
    dashboard = wb.active
    dashboard.title = "Dashboard"
    build_dashboard(dashboard, meta, stage_results, merged, flow_findings,
                    len(control_rows), hotspots, git_summary, warnings)

    def cwe(tags):
        found = {m.upper() for t in tags
                 for m in re.findall(r"cwe-\d+", t.lower())}
        return ", ".join(sorted(found))

    finding_headers = [
        "ID", "Severity", "Security Score", "CWE", "Rule ID", "Rule Name",
        "File", "Line", "Message", "Reported By", "Flow Path?", "Description",
    ]

    def finding_row(f):
        return [
            f["id"], f["severity"],
            "" if f["score"] is None else f["score"], cwe(f["tags"]),
            f["rule_id"], f["rule_name"], f["file"], f["line"],
            f["message"], "\n".join(sorted(f["sources"])),
            "Yes" if f["flows"] else "No", f["description"],
        ]

    finding_widths = {
        "ID": 9, "Severity": 12, "Security Score": 10, "CWE": 14,
        "Rule ID": 30, "Rule Name": 28, "File": 38, "Line": 7,
        "Message": 60, "Reported By": 34, "Flow Path?": 10,
        "Description": 45,
    }
    wrap = ("Message", "Reported By", "Description", "Rule Name", "File")

    security = [
        f for f in merged
        if f["score"] is not None
        or any("security" in t.lower() or "cwe" in t.lower()
               for t in f["tags"])
    ]
    write_table(wb, "Security Findings", finding_headers,
                [finding_row(f) for f in security], widths=finding_widths,
                wrap=wrap, styled=("Severity",),
                empty_note="No security-tagged findings were reported.")
    write_table(wb, "All Findings", finding_headers,
                [finding_row(f) for f in merged], widths=finding_widths,
                wrap=wrap, styled=("Severity",),
                empty_note="No findings were reported.")

    # rule summary
    by_rule: dict[str, list[dict]] = defaultdict(list)
    for f in merged:
        by_rule[f["rule_id"]].append(f)
    rule_rows = []
    for rule_id, items in by_rule.items():
        worst = max(items, key=lambda f: SEV_RANK[f["severity"]])
        rule_rows.append([
            rule_id, worst["rule_name"], worst["severity"], len(items),
            len({i["file"] for i in items}),
            "Yes" if any(i["score"] is not None for i in items) else "No",
            worst["description"],
        ])
    rule_rows.sort(key=lambda r: (-SEV_RANK[r[2]], -r[3], r[0]))
    write_table(
        wb, "Rule Summary",
        ["Rule ID", "Rule Name", "Worst Severity", "Findings",
         "Files Affected", "Security Rule?", "Description"],
        rule_rows,
        widths={"Rule ID": 34, "Rule Name": 30, "Description": 60,
                "Worst Severity": 14},
        wrap=("Rule Name", "Description"), styled=("Worst Severity",),
    )

    write_table(
        wb, "File Hotspots",
        ["File", "Findings", "Critical", "High", "Medium", "Low", "Top Rule",
         "AST Prompts", "Lines", "Functions", "Long Functions",
         "Bare Excepts", "Git Status", "Git Priority"],
        hotspots,
        widths={"File": 44, "Top Rule": 32},
        wrap=("File",), styled=("Git Priority",),
    )

    write_table(
        wb, "Data Flow Paths",
        ["Finding ID", "Rule ID", "Severity", "Alert Location", "Path #",
         "Step #", "Step Message", "File", "Line", "Code Snippet"],
        flow_rows,
        widths={"Rule ID": 30, "Alert Location": 34, "Step Message": 50,
                "File": 38, "Code Snippet": 50, "Finding ID": 10,
                "Severity": 12, "Path #": 8, "Step #": 8, "Line": 7},
        wrap=("Step Message", "Code Snippet", "File", "Alert Location"),
        styled=("Severity",),
        empty_note="CodeQL supplied no explicit flow paths in this run.",
    )

    write_table(
        wb, "Control Flow Candidates",
        ["Severity", "Rule ID", "Rule Name", "File", "Line", "Message",
         "Matched Keywords"],
        control_rows,
        widths={"Rule ID": 30, "Rule Name": 28, "File": 38, "Message": 60,
                "Matched Keywords": 28, "Severity": 12, "Line": 7},
        wrap=("Message", "File", "Rule Name", "Matched Keywords"),
        styled=("Severity",),
        empty_note="No findings matched the control-flow keywords.",
    )

    git_headers = (changed_rows[0] if changed_rows else
                   ["Working Status", "Diff Status", "Path", "Previous Path",
                    "Priority", "Reason"])
    write_table(wb, "Git Changes", [str(h) for h in git_headers],
                changed_rows[1:], widths={"Path": 50, "Reason": 55},
                wrap=("Path", "Reason", "Previous Path"),
                styled=("Priority",),
                empty_note="No changes vs baseline (or change audit did not run).")

    custom = read_sheet(change_xlsx, "Custom Queries")
    write_table(wb, "Custom Queries",
                [str(h) for h in (custom[0] if custom else
                 ["Query File", "Status", "Path/SARIF", "Exit Code", "Notes"])],
                custom[1:],
                widths={"Path/SARIF": 55, "Notes": 60},
                wrap=("Path/SARIF", "Notes"),
                empty_note="Custom queries were not part of this run.")
    custom_findings = read_sheet(change_xlsx, "Custom Query Findings")
    if len(custom_findings) > 1:
        write_table(wb, "Custom Query Findings",
                    [str(h) for h in custom_findings[0]], custom_findings[1:],
                    wrap=("Message",), widths={"Message": 60})

    write_table(wb, "AST Metrics",
                [str(h) for h in (read_sheet(quality_xlsx, "File Metrics") or
                 [["File"]])[0]],
                read_sheet(quality_xlsx, "File Metrics")[1:],
                widths={"File": 50}, wrap=("File",),
                empty_note="quality_audit did not produce metrics.")
    write_table(wb, "AST Review",
                [str(h) for h in (ast_rows or [["File"]])[0]], ast_rows[1:],
                widths={"File": 44, "Detail": 70}, wrap=("File", "Detail"),
                styled=("Severity",),
                empty_note="No AST review candidates.")

    write_table(
        wb, "Run Log",
        ["Stage", "Status", "Seconds", "Command", "Details",
         "Last Output Lines"],
        [[s.title, s.status, s.seconds, s.command, s.detail,
          "\n".join(s.tail[-25:])] for s in stage_results],
        widths={"Stage": 36, "Command": 60, "Details": 36,
                "Last Output Lines": 90},
        wrap=("Stage", "Command", "Details", "Last Output Lines"),
        styled=("Status",),
    )

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)

    return {
        "unique": len(merged),
        "security": len(security),
        "flow_steps": len(flow_rows),
        "control": len(control_rows),
        "files": len(hotspots),
    }


# --------------------------------------------------------------------------
# Orchestration
# --------------------------------------------------------------------------

def parse_args() -> argparse.Namespace:
    ap = argparse.ArgumentParser(
        description="Run all Mythri audits and write one Excel report."
    )
    ap.add_argument("--repo", type=Path, default=DEFAULT_REPO,
                    help=f"Repository root (default: {DEFAULT_REPO})")
    ap.add_argument("--out", type=Path, default=DEFAULT_OUT,
                    help=f"Output root (default: {DEFAULT_OUT})")
    ap.add_argument("--skip", nargs="+", choices=STAGE_KEYS, default=[],
                    metavar="STAGE",
                    help="Stages to skip: " + ", ".join(STAGE_KEYS))
    ap.add_argument("--baseline", default="HEAD",
                    help="Git revision for the change audit (default HEAD)")
    ap.add_argument("--with-custom-queries", action="store_true",
                    help="Run the custom .ql templates in the change audit")
    ap.add_argument("--quality-skip-codeql", action="store_true",
                    help="Run only AST checks inside quality_audit.py")
    return ap.parse_args()


def main() -> int:
    args = parse_args()
    repo = args.repo.expanduser().resolve()

    if not (repo / ".git").exists():
        print(f"ERROR: {repo} is not a Git repository.\n"
              "Use --repo /path/to/mythri_slm", file=sys.stderr)
        return 2

    missing = [n for n in ("codeql_audit.py", "data_flow_audit.py",
                           "control_flow_audit.py", "quality_audit.py",
                           "mythri_change_audit.py")
               if not (SCRIPTS_DIR / n).is_file()]
    if missing:
        print(f"ERROR: missing scripts in {SCRIPTS_DIR}: {', '.join(missing)}",
              file=sys.stderr)
        return 2

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    run_dir = args.out.expanduser().resolve() / stamp
    run_dir.mkdir(parents=True, exist_ok=False)

    py = sys.executable
    cli = find_codeql()
    suite_dir = find_suite_dir()
    started_at = time.time()

    print("MYTHRI MASTER AUDIT")
    print(f"Repository : {repo}")
    print(f"Output     : {run_dir}")
    print(f"CodeQL CLI : {cli or 'NOT FOUND'}")
    print(f"Query suite: {suite_dir or 'NOT FOUND'}")

    results: list[StageResult] = []
    db_failed = False

    def codeql_missing() -> str:
        if not cli:
            return "CodeQL CLI not found"
        if not suite_dir:
            return "Python query pack not found"
        if not DB_PATH.exists():
            return f"CodeQL database not found: {DB_PATH}"
        return ""

    # 1 ---------------------------------------------------------------
    title = "1. CodeQL database rebuild + 3 suites (codeql_audit.py)"
    if "codeql" in args.skip:
        results.append(skipped("codeql", title, "skipped by --skip"))
    else:
        res = run_stage("codeql", title, [py, SCRIPTS_DIR / "codeql_audit.py"],
                        repo)
        db_failed = res.status != "SUCCESS"
        results.append(res)

    # 2 ---------------------------------------------------------------
    title = "2. Data-flow audit (data_flow_audit.py)"
    problem = codeql_missing()
    if "data_flow" in args.skip:
        results.append(skipped("data_flow", title, "skipped by --skip"))
    elif db_failed:
        results.append(skipped("data_flow", title,
                               "database rebuild failed; results would be stale"))
    elif problem:
        results.append(skipped("data_flow", title, problem))
    else:
        results.append(run_stage("data_flow", title, [
            py, SCRIPTS_DIR / "data_flow_audit.py",
            "--codeql", cli, "--database", DB_PATH,
            "--suite", suite_dir / "python-security-extended.qls",
            "--output-dir", run_dir / "data_flow",
        ], repo))

    # 3 ---------------------------------------------------------------
    title = "3. Control-flow audit (control_flow_audit.py)"
    if "control_flow" in args.skip:
        results.append(skipped("control_flow", title, "skipped by --skip"))
    elif db_failed:
        results.append(skipped("control_flow", title,
                               "database rebuild failed; results would be stale"))
    elif problem:
        results.append(skipped("control_flow", title, problem))
    else:
        results.append(run_stage("control_flow", title, [
            py, SCRIPTS_DIR / "control_flow_audit.py",
            "--codeql", cli, "--database", DB_PATH,
            "--suite", suite_dir / "python-code-quality-extended.qls",
            "--output", run_dir / "control_flow",
        ], repo))

    # 4 ---------------------------------------------------------------
    title = "4. Code quality audit (quality_audit.py)"
    if "quality" in args.skip:
        results.append(skipped("quality", title, "skipped by --skip"))
    else:
        cmd = [py, SCRIPTS_DIR / "quality_audit.py",
               "--repo", repo, "--out", run_dir / "quality"]
        if args.quality_skip_codeql or db_failed:
            cmd.append("--skip-codeql")
        results.append(run_stage("quality", title, cmd, repo))

    # 5 ---------------------------------------------------------------
    title = "5. Git change audit (mythri_change_audit.py)"
    if "changes" in args.skip:
        results.append(skipped("changes", title, "skipped by --skip"))
    else:
        cmd = [py, SCRIPTS_DIR / "mythri_change_audit.py",
               "--repo", repo, "--out", run_dir / "change",
               "--baseline", args.baseline]
        if cli:
            cmd += ["--codeql", cli]
        cmd += ["--database", DB_PATH]
        if args.with_custom_queries:
            cmd.append("--run-codeql")
        results.append(run_stage("changes", title, cmd, repo))

    # ---- collect outputs --------------------------------------------
    warnings: list[str] = []
    parsed: list[dict[str, Any]] = []

    def fresh(path: Path) -> bool:
        return path.is_file() and path.stat().st_mtime >= started_at - 5

    def ingest(path: Path, label: str) -> None:
        try:
            parsed.extend(parse_sarif(path, label))
        except (OSError, json.JSONDecodeError) as exc:
            warnings.append(f"Could not read {path.name}: {exc}")

    for suite in LEGACY_SUITES:
        path = LEGACY_REPORT_DIR / f"{suite}.sarif"
        if fresh(path):
            ingest(path, f"codeql_audit: {suite}")

    for path in sorted((run_dir / "data_flow").glob("**/data_flow_*.sarif")):
        ingest(path, "data_flow_audit: security-extended")

    control_sarifs = sorted(
        (run_dir / "control_flow").glob("**/control_flow_full.sarif"))
    for path in control_sarifs:
        ingest(path, "control_flow_audit: quality-extended")

    for path in sorted((run_dir / "quality").glob("mythri_quality_*.sarif")):
        ingest(path, "quality_audit: quality-extended")

    quality_books = sorted((run_dir / "quality").glob(
        "mythri_quality_audit_*.xlsx"))
    quality_xlsx = quality_books[-1] if quality_books else None
    change_books = sorted((run_dir / "change").glob(
        "mythri_change_audit_*.xlsx"))
    change_xlsx = change_books[-1] if change_books else None

    # quality_audit exits 0 even if its CodeQL step failed; surface that.
    for row in read_sheet(quality_xlsx, "Summary"):
        if row and row[0] == "CodeQL status":
            for res in results:
                if res.key == "quality" and res.status == "SUCCESS":
                    res.detail = f"CodeQL step: {row[1]}"

    control_rows, control_warning = load_control_flow_rows(
        control_sarifs[-1] if control_sarifs else None)
    if control_warning and "control_flow" not in args.skip:
        warnings.append(control_warning)

    if not parsed:
        warnings.append(
            "No CodeQL findings were collected. Check the Run Log sheet to "
            "see whether the CodeQL stages ran.")
    if "codeql" in args.skip and DB_PATH.exists():
        warnings.append(
            "Database rebuild was skipped; findings reflect the database as "
            "of the 'Database last modified' time, not necessarily the "
            "current source.")

    db_mtime = (datetime.fromtimestamp(DB_PATH.stat().st_mtime)
                .astimezone().isoformat(timespec="seconds")
                if DB_PATH.exists() else "database not found")

    meta = {
        "generated": datetime.now().astimezone().isoformat(timespec="seconds"),
        "repo": str(repo),
        "run_dir": str(run_dir),
        "cli": str(cli or "not found"),
        "db": str(DB_PATH),
        "db_mtime": db_mtime,
    }

    out_path = run_dir / f"mythri_master_audit_{stamp}.xlsx"
    counts = build_report(out_path, meta, results, parsed, control_rows,
                          quality_xlsx, change_xlsx, warnings)

    print(f"\n{'=' * 76}\nMASTER AUDIT COMPLETE\n{'=' * 76}")
    for res in results:
        print(f"  [{res.status:<7}] {res.title}")
    print(f"\nUnique findings        : {counts['unique']}")
    print(f"Security findings      : {counts['security']}")
    print(f"Data-flow steps        : {counts['flow_steps']}")
    print(f"Control-flow candidates: {counts['control']}")
    print(f"Files in hotspot sheet : {counts['files']}")
    print(f"\nExcel report : {out_path}")
    print(f"Raw outputs  : {run_dir}")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except KeyboardInterrupt:
        print("\nInterrupted by user.", file=sys.stderr)
        raise SystemExit(130)